import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# PREMIUM EDITORIAL STUDIO LIGHT THEME FOR PPT
# ==============================================================================
BG_SLIDE = RGBColor(248, 250, 252)      # #F8FAFC - Clean Architectural Off-White Canvas
SURFACE = RGBColor(255, 255, 255)       # #FFFFFF - Crisp Pure White Card
SURFACE_LIGHT = RGBColor(241, 245, 249) # #F1F5F9 - Subtle Neutral Container
SURFACE_TINT = RGBColor(238, 242, 255)  # #EEF2FF - Soft Indigo Tint
BORDER_SUBTLE = RGBColor(226, 232, 240) # #E2E8F0 - Refined Hairline Divider
BORDER_MUTED = RGBColor(203, 213, 225)  # #CBD5E1 - Structural Card Border
BORDER_ACTIVE = RGBColor(147, 197, 253) # #93C5FD - Active highlight

ACCENT_BLUE = RGBColor(37, 99, 235)     # #2563EB - Royal Tech Blue
ACCENT_INDIGO = RGBColor(79, 70, 229)   # #4F46E5 - Electric Indigo
ACCENT_EMERALD = RGBColor(5, 150, 105)  # #059669 - Forest Mint
ACCENT_PURPLE = RGBColor(124, 58, 237)  # #7C3AED - Royal Violet

TEXT_PRIMARY = RGBColor(15, 23, 42)     # #0F172A - Deep Charcoal Ink (Maximum Contrast)
TEXT_SECONDARY = RGBColor(51, 65, 85)   # #334155 - Slate Grey Body
TEXT_MUTED = RGBColor(100, 116, 139)    # #64748B - Subtle Caption

SCREENSHOTS_DIR = os.path.join(os.path.dirname(__file__), '..', 'presentation', 'screenshots')
OUTPUT_PATH = os.path.join(os.path.dirname(__file__), '..', 'presentation', 'Swetha_Portfolio_Presentation.pptx')

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def set_slide_background(slide, color=BG_SLIDE):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_eyebrow(slide, text, top=Inches(0.65), left=Inches(0.9)):
    box = slide.shapes.add_textbox(left, top, Inches(8.0), Inches(0.35))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text.upper()
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO
    p.font.name = "Outfit"

def add_header(slide, eyebrow, title, subtitle=None, top_offset=Inches(0.65)):
    # Eyebrow
    add_eyebrow(slide, eyebrow, top=top_offset)
    
    # Title
    tbox = slide.shapes.add_textbox(Inches(0.9), top_offset + Inches(0.32), Inches(11.5), Inches(0.7))
    tf = tbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = TEXT_PRIMARY
    p.font.name = "Outfit"
    
    if subtitle:
        sbox = slide.shapes.add_textbox(Inches(0.9), top_offset + Inches(1.0), Inches(11.5), Inches(0.4))
        stf = sbox.text_frame
        stf.word_wrap = True
        sp = stf.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(13)
        sp.font.color.rgb = TEXT_SECONDARY
        sp.font.name = "Inter"

# ==============================================================================
# SLIDE 1: Striking Editorial Title Slide (Clean Keynote Layout, No Screenshot)
# ==============================================================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1)

# Main Hero Card (Expansive, Elegant Studio Frame)
hero_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.85), Inches(11.533), Inches(5.8))
hero_card.fill.solid()
hero_card.fill.fore_color.rgb = SURFACE
hero_card.line.color.rgb = BORDER_SUBTLE
hero_card.line.width = Pt(1)

# Eyebrow Pill
eye_pill = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.3), Inches(4.5), Inches(0.42))
eye_pill.fill.solid()
eye_pill.fill.fore_color.rgb = SURFACE_TINT
eye_pill.line.color.rgb = BORDER_ACTIVE
eye_pill.line.width = Pt(1)
p_pill = eye_pill.text_frame.paragraphs[0]
p_pill.text = "ACADEMIC CAPSTONE & TECHNICAL EVALUATION"
p_pill.font.size = Pt(9.5)
p_pill.font.bold = True
p_pill.font.color.rgb = ACCENT_INDIGO
p_pill.font.name = "JetBrains Mono"
p_pill.alignment = PP_ALIGN.CENTER

# Main Title
tb_t = slide1.shapes.add_textbox(Inches(1.5), Inches(1.95), Inches(10.333), Inches(1.6))
tf_t = tb_t.text_frame
tf_t.word_wrap = True
p = tf_t.paragraphs[0]
p.text = "Personal Developer Portfolio Website"
p.font.size = Pt(38)
p.font.bold = True
p.font.color.rgb = TEXT_PRIMARY
p.font.name = "Outfit"

# Subtitle
tb_sub = slide1.shapes.add_textbox(Inches(1.5), Inches(2.95), Inches(9.8), Inches(0.8))
tf_sub = tb_sub.text_frame
tf_sub.word_wrap = True
p = tf_sub.paragraphs[0]
p.text = "A production-grade digital platform engineered with semantic web standards, applied RAG intelligence, and human-centered design principles."
p.font.size = Pt(14)
p.font.color.rgb = TEXT_SECONDARY
p.font.name = "Inter"

# Divider Line
div = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(3.9), Inches(10.333), Inches(0.02))
div.fill.solid()
div.fill.fore_color.rgb = BORDER_SUBTLE
div.line.fill.background()

# Bottom 4-Column Metadata Grid
meta_columns = [
    ("PRESENTER", "Swetha N.", "B.E. Computer Science and Design", ACCENT_INDIGO),
    ("INSTITUTION", "Rajalakshmi Eng. College", "Chennai · Class of 2028 (CGPA: 8.65)", ACCENT_BLUE),
    ("CORE DOMAINS", "Backend & Applied AI", "RAG · Spring Boot · Django · Systems", ACCENT_EMERALD),
    ("LIVE PLATFORM", "Verified Production CDN", "my-portfolio-...vercel.app", ACCENT_PURPLE)
]

for i, (tag, val1, val2, col) in enumerate(meta_columns):
    mx = Inches(1.5 + i * 2.65)
    my = Inches(4.15)
    
    m_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mx, my, Inches(2.45), Inches(2.0))
    m_box.fill.solid()
    m_box.fill.fore_color.rgb = SURFACE_LIGHT
    m_box.line.color.rgb = BORDER_SUBTLE
    m_box.line.width = Pt(1)
    
    # Colored top indicator
    m_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, mx, my, Inches(2.45), Inches(0.04))
    m_bar.fill.solid()
    m_bar.fill.fore_color.rgb = col
    m_bar.line.fill.background()
    
    tf_m = m_box.text_frame
    tf_m.word_wrap = True
    
    p_tag = tf_m.paragraphs[0]
    p_tag.text = tag
    p_tag.font.size = Pt(9)
    p_tag.font.bold = True
    p_tag.font.color.rgb = col
    p_tag.font.name = "JetBrains Mono"
    
    p_v1 = tf_m.add_paragraph()
    p_v1.text = val1
    p_v1.font.size = Pt(13)
    p_v1.font.bold = True
    p_v1.font.color.rgb = TEXT_PRIMARY
    p_v1.font.name = "Outfit"
    p_v1.space_before = Pt(6)
    
    p_v2 = tf_m.add_paragraph()
    p_v2.text = val2
    p_v2.font.size = Pt(9.5)
    p_v2.font.color.rgb = TEXT_MUTED
    p_v2.font.name = "Inter"
    p_v2.space_before = Pt(4)

# ==============================================================================
# SLIDE 2: The Idea & Objective (5-Second Visual Hierarchy)
# ==============================================================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2)
add_header(slide2, "Core Philosophy", "The Problem & The Purpose", 
           "Moving beyond flat resumes to establish verified engineering credibility.")

# Top Half: High-Impact Problem Statement Card
prob_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.9), Inches(11.533), Inches(1.6))
prob_card.fill.solid()
prob_card.fill.fore_color.rgb = SURFACE_TINT
prob_card.line.color.rgb = ACCENT_INDIGO
prob_card.line.width = Pt(1.5)

tf_p = prob_card.text_frame
tf_p.word_wrap = True
p = tf_p.paragraphs[0]
p.text = "THE PROBLEM"
p.font.size = Pt(10)
p.font.bold = True
p.font.color.rgb = ACCENT_INDIGO
p.font.name = "Outfit"

p2 = tf_p.add_paragraph()
p2.text = "Static PDF resumes fail to convey live engineering craftsmanship, interactive problem-solving, or system design."
p2.font.size = Pt(16)
p2.font.bold = True
p2.font.color.rgb = TEXT_PRIMARY
p2.font.name = "Outfit"
p2.space_before = Pt(4)

p3 = tf_p.add_paragraph()
p3.text = "Developers need an authentic digital platform that proves capability through live code, production internship outcomes, and architectural depth."
p3.font.size = Pt(12)
p3.font.color.rgb = TEXT_SECONDARY
p3.font.name = "Inter"
p3.space_before = Pt(4)

# Bottom Half: 3 Clear, Bold Objective Cards (Horizontal Columns)
objectives_data = [
    ("01", "Showcase Production Engineering", 
     "Highlight genuine enterprise experience from Zoho Corporation — specifically OpenAPI RAG workflows, vector embeddings, and backend services.", ACCENT_INDIGO),
    ("02", "Prove Algorithmic Fluency", 
     "Provide verified credentials through 500+ LeetCode problems solved across Trees, Graphs, Dynamic Programming, and Competitive Hackathons.", ACCENT_EMERALD),
    ("03", "Deliver Thoughtful Ergonomics", 
     "Execute a zero-framework, accessible web experience (<50ms load time) with persistent dark/light theming and human-centered design.", ACCENT_BLUE)
]

for i, (num, heading, detail, color) in enumerate(objectives_data):
    x = Inches(0.9 + i * 3.98)
    y = Inches(3.8)
    
    card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.6), Inches(2.8))
    card.fill.solid()
    card.fill.fore_color.rgb = SURFACE
    card.line.color.rgb = BORDER_SUBTLE
    card.line.width = Pt(1)
    
    # Top color line
    bar = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(3.6), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    
    tf = card.text_frame
    tf.word_wrap = True
    
    p_num = tf.paragraphs[0]
    p_num.text = num
    p_num.font.size = Pt(18)
    p_num.font.bold = True
    p_num.font.color.rgb = color
    p_num.font.name = "JetBrains Mono"
    
    p_h = tf.add_paragraph()
    p_h.text = heading
    p_h.font.size = Pt(14)
    p_h.font.bold = True
    p_h.font.color.rgb = TEXT_PRIMARY
    p_h.font.name = "Outfit"
    p_h.space_before = Pt(6)
    
    p_d = tf.add_paragraph()
    p_d.text = detail
    p_d.font.size = Pt(11)
    p_d.font.color.rgb = TEXT_SECONDARY
    p_d.font.name = "Inter"
    p_d.space_before = Pt(8)

# ==============================================================================
# SLIDE 3: Website Structure & User Journey (Visual Flow Diagram)
# ==============================================================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3)
add_header(slide3, "User Journey & Experience", "Website Structure & Visitor Flow", 
           "How a technical reviewer or recruiter moves through the portfolio seamlessly.")

# Top: Visual Journey Steps Flow (7 interconnected nodes)
journey_steps = [
    ("01", "Hero", "First Impression & Role"),
    ("02", "About", "Academic & CS Philosophy"),
    ("03", "Skills", "6 Concrete Domains"),
    ("04", "Projects", "7 Verified Case Studies"),
    ("05", "Experience", "Zoho SWE & LeetCode"),
    ("06", "Resume", "PDF & Summary Badges"),
    ("07", "Contact", "Validated Recruiter Form")
]

for i, (num, step_name, desc) in enumerate(journey_steps):
    x = Inches(0.9 + i * 1.68)
    y = Inches(1.9)
    node = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(1.5), Inches(1.6))
    node.fill.solid()
    node.fill.fore_color.rgb = SURFACE
    node.line.color.rgb = ACCENT_INDIGO if i in [0, 3, 4] else BORDER_SUBTLE
    node.line.width = Pt(1.2)
    
    tf = node.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.font.name = "JetBrains Mono"
    
    p2 = tf.add_paragraph()
    p2.text = step_name
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_PRIMARY
    p2.font.name = "Outfit"
    p2.space_before = Pt(4)
    
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(8.5)
    p3.font.color.rgb = TEXT_SECONDARY
    p3.font.name = "Inter"
    p3.space_before = Pt(4)

# Bottom: 3 Key UX Decisions (Generous Cards)
ux_decisions = [
    ("Editorial Dark Theme & Zero Flash", 
     "A refined midnight palette (#090D16) paired with an inline localStorage theme-engine prevents white flashes on reload while offering dark/light comfort.", ACCENT_BLUE),
    ("Evidence-Based Competencies", 
     "Replaced arbitrary skill percentage bars with categorized, tangible toolsets and verifiable milestone counts (500+ LeetCode problems solved).", ACCENT_INDIGO),
    ("Frictionless Recruiter Conversion", 
     "Quick-view modals allow rapid project inspection without leaving context. One-click resume download and client-validated contact channels facilitate direct outreach.", ACCENT_EMERALD)
]

for i, (title, text, color) in enumerate(ux_decisions):
    x = Inches(0.9 + i * 3.98)
    y = Inches(3.9)
    card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.6), Inches(2.7))
    card.fill.solid()
    card.fill.fore_color.rgb = SURFACE
    card.line.color.rgb = BORDER_SUBTLE
    card.line.width = Pt(1)
    
    bar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(3.6), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    
    tf = card.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_PRIMARY
    p.font.name = "Outfit"
    
    p2 = tf.add_paragraph()
    p2.text = text
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_SECONDARY
    p2.font.name = "Inter"
    p2.space_before = Pt(10)

# ==============================================================================
# SLIDE 4: Technology & Implementation (Clean 4-Quadrant Architecture)
# ==============================================================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4)
add_header(slide4, "Engineering Stack", "Technology & Implementation Flow", 
           "A disciplined, zero-framework stack prioritizing speed, accessibility, and reliability.")

tech_quads = [
    ("01 · Structure", "Semantic HTML5", ACCENT_EMERALD, [
        "ARIA landmark architecture (<header>, <main>, <nav>, <section>)",
        "Accessibility-first keyboard skip-link navigation",
        "OpenGraph & Twitter meta tags for rich social indexing"
    ]),
    ("02 · Styling", "Vanilla CSS3 System", ACCENT_BLUE, [
        "CSS Custom Properties (:root) for instant theme variables",
        "Modern CSS Grid & Flexbox for multi-device responsiveness",
        "Subtle glassmorphic backdrops and fluid clamp() typography"
    ]),
    ("03 · Interactivity", "Vanilla JavaScript (ES6+)", ACCENT_INDIGO, [
        "Zero external framework bloat — pure browser native code",
        "Theme persistence controller with localStorage synchronization",
        "Active scroll-spy navigation listener & Quick View modal logic"
    ]),
    ("04 · Infrastructure", "Git, GitHub & Vercel", ACCENT_PURPLE, [
        "Production version control on GitHub main branch",
        "Automated continuous deployment via Vercel Edge CDN",
        "Zero build-step pipeline ensuring instantaneous deployments"
    ])
]

for i, (quad_label, quad_title, col, pts) in enumerate(tech_quads):
    x = Inches(0.9 + (i % 2) * 5.95)
    y = Inches(1.9 + (i // 2) * 2.5)
    
    card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.25))
    card.fill.solid()
    card.fill.fore_color.rgb = SURFACE
    card.line.color.rgb = BORDER_SUBTLE
    card.line.width = Pt(1)
    
    # Left accent indicator
    bar = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.08), Inches(2.25))
    bar.fill.solid()
    bar.fill.fore_color.rgb = col
    bar.line.fill.background()
    
    tf = card.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = quad_label.upper()
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = col
    p.font.name = "JetBrains Mono"
    
    p_t = tf.add_paragraph()
    p_t.text = quad_title
    p_t.font.size = Pt(16)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_PRIMARY
    p_t.font.name = "Outfit"
    p_t.space_before = Pt(2)
    p_t.space_after = Pt(6)
    
    for pt in pts:
        p_pt = tf.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.size = Pt(10.5)
        p_pt.font.color.rgb = TEXT_SECONDARY
        p_pt.font.name = "Inter"
        p_pt.space_after = Pt(3)

# ==============================================================================
# SLIDE 5: The Actual Website (Large, Prominent Real Screenshots)
# ==============================================================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5)
add_header(slide5, "Visual Proof", "The Implemented Website", 
           "Genuine high-resolution captures showing the live design system in action.")

# We showcase 2 large, prominent screenshots side by side: Hero Showcase & Featured Projects Showcase
showcases = [
    ("01_hero.png", "Hero & Selected Work Overview", 
     "Balanced personal introduction, role pill, live status indicator, and multi-project quick cards.", Inches(0.9)),
    ("04_projects.png", "Featured Projects Showcase", 
     "Deep engineering breakdowns: Zoho OpenAPI RAG Assistant, Healthcare EHR, and Smart SMS Django Platform.", Inches(6.8))
]

for filename, title, desc, x in showcases:
    card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.9), Inches(5.633), Inches(5.0))
    card.fill.solid()
    card.fill.fore_color.rgb = SURFACE
    card.line.color.rgb = BORDER_SUBTLE
    card.line.width = Pt(1)
    
    img_path = os.path.join(SCREENSHOTS_DIR, filename)
    if os.path.exists(img_path):
        slide5.shapes.add_picture(img_path, x + Inches(0.15), Inches(2.05), Inches(5.333), Inches(3.35))
    
    # Caption Area
    tb = slide5.shapes.add_textbox(x + Inches(0.2), Inches(5.5), Inches(5.2), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.font.name = "Outfit"
    p.space_after = Pt(4)
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = TEXT_SECONDARY
    p2.font.name = "Inter"

# ==============================================================================
# SLIDE 6: Outcome & Conclusion (Confident, Minimal Editorial Finish)
# ==============================================================================
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6)
add_header(slide6, "Project Summary", "Outcomes, Roadmap & Conclusion", 
           "Verified deliverables, technical evaluation, and future engineering trajectory.")

# Left Card: Verified Outcomes (Key Achievements)
outcomes_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.9), Inches(5.6), Inches(4.8))
outcomes_card.fill.solid()
outcomes_card.fill.fore_color.rgb = SURFACE
outcomes_card.line.color.rgb = BORDER_SUBTLE
outcomes_card.line.width = Pt(1)

bar_o = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.9), Inches(5.6), Inches(0.06))
bar_o.fill.solid()
bar_o.fill.fore_color.rgb = ACCENT_EMERALD
bar_o.line.fill.background()

tf_o = outcomes_card.text_frame
tf_o.word_wrap = True
p = tf_o.paragraphs[0]
p.text = "VERIFIED PROJECT OUTCOMES"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_EMERALD
p.font.name = "JetBrains Mono"

points_out = [
    ("Production-Ready Deployment", "Successfully hosted on Vercel Edge with zero downtime, instant SSL, and worldwide CDN distribution."),
    ("Zero-Dependency Architecture", "Pure HTML5, CSS3, and ES6+ delivering sub-50ms render speed without heavy client bundle overhead."),
    ("Dual-Theme Precision", "Flawless dark and light modes with localStorage persistence and zero layout shift."),
    ("Authentic Technical Depth", "Seamlessly integrates real Zoho Corporation RAG internship contributions and 500+ LeetCode milestones.")
]

for title, desc in points_out:
    p_t = tf_o.add_paragraph()
    p_t.text = "✓ " + title
    p_t.font.size = Pt(13)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_PRIMARY
    p_t.font.name = "Outfit"
    p_t.space_before = Pt(8)
    
    p_d = tf_o.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(10.5)
    p_d.font.color.rgb = TEXT_SECONDARY
    p_d.font.name = "Inter"

# Right Card: Future Scope & Roadmap
future_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.633), Inches(4.8))
future_card.fill.solid()
future_card.fill.fore_color.rgb = SURFACE
future_card.line.color.rgb = BORDER_SUBTLE
future_card.line.width = Pt(1)

bar_f = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.633), Inches(0.06))
bar_f.fill.solid()
bar_f.fill.fore_color.rgb = ACCENT_INDIGO
bar_f.line.fill.background()

tf_f = future_card.text_frame
tf_f.word_wrap = True

p = tf_f.paragraphs[0]
p.text = "FUTURE ENGINEERING ROADMAP"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_INDIGO
p.font.name = "JetBrains Mono"

scope_pts = [
    ("Live LeetCode & GitHub Sync", "Dynamic automated pipeline refreshing solved DSA milestones and live repo activity."),
    ("Interactive Semantic RAG Playground", "In-browser live demonstration of OpenAPI and YAML semantic vector retrieval with LangChain."),
    ("Automated Lighthouse CI Gates", "Continuous integration checks enforcing 98+ accessibility, performance, and SEO scores on commits."),
    ("Ambient WebGL Micro-Interactions", "Subtle GPU-accelerated interactive canvas shaders for featured project case studies.")
]

for title, desc in scope_pts:
    p_t = tf_f.add_paragraph()
    p_t.text = "→ " + title
    p_t.font.size = Pt(13)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_PRIMARY
    p_t.font.name = "Outfit"
    p_t.space_before = Pt(8)
    
    p_d = tf_f.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(10.5)
    p_d.font.color.rgb = TEXT_SECONDARY
    p_d.font.name = "Inter"

# ==============================================================================
# SLIDE 7: Dedicated Professional Thank You Slide
# ==============================================================================
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7)

# Main Hero Card (Expansive, Elegant Keynote Frame)
thx_card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.85), Inches(11.533), Inches(5.8))
thx_card.fill.solid()
thx_card.fill.fore_color.rgb = SURFACE
thx_card.line.color.rgb = BORDER_SUBTLE
thx_card.line.width = Pt(1)

# Eyebrow Pill
thx_pill = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.416), Inches(1.25), Inches(4.5), Inches(0.42))
thx_pill.fill.solid()
thx_pill.fill.fore_color.rgb = SURFACE_TINT
thx_pill.line.color.rgb = BORDER_ACTIVE
thx_pill.line.width = Pt(1)
p_tpill = thx_pill.text_frame.paragraphs[0]
p_tpill.text = "CAPSTONE EVALUATION & CONCLUSION"
p_tpill.font.size = Pt(9.5)
p_tpill.font.bold = True
p_tpill.font.color.rgb = ACCENT_INDIGO
p_tpill.font.name = "JetBrains Mono"
p_tpill.alignment = PP_ALIGN.CENTER

# Main Title: "Thank You!"
tb_thx_title = slide7.shapes.add_textbox(Inches(1.5), Inches(1.85), Inches(10.333), Inches(1.1))
tf_tt = tb_thx_title.text_frame
tf_tt.word_wrap = True
p_tt = tf_tt.paragraphs[0]
p_tt.text = "Thank You!"
p_tt.font.size = Pt(44)
p_tt.font.bold = True
p_tt.font.color.rgb = TEXT_PRIMARY
p_tt.font.name = "Outfit"
p_tt.alignment = PP_ALIGN.CENTER

# Subtitle / Q&A Prompt
tb_thx_sub = slide7.shapes.add_textbox(Inches(1.5), Inches(2.9), Inches(10.333), Inches(0.6))
tf_ts = tb_thx_sub.text_frame
tf_ts.word_wrap = True
p_ts = tf_ts.paragraphs[0]
p_ts.text = "Open for Technical Questions, Architectural Discussions & Feedback"
p_ts.font.size = Pt(15)
p_ts.font.bold = False
p_ts.font.color.rgb = TEXT_SECONDARY
p_ts.font.name = "Inter"
p_ts.alignment = PP_ALIGN.CENTER

# Divider Line
div_thx = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.5), Inches(3.6), Inches(8.333), Inches(0.02))
div_thx.fill.solid()
div_thx.fill.fore_color.rgb = BORDER_SUBTLE
div_thx.line.fill.background()

# Presenter Info Banner / Box
p_banner = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(3.85), Inches(10.333), Inches(0.95))
p_banner.fill.solid()
p_banner.fill.fore_color.rgb = SURFACE_LIGHT
p_banner.line.color.rgb = BORDER_SUBTLE
p_banner.line.width = Pt(1)

tf_pb = p_banner.text_frame
tf_pb.word_wrap = True
p_pb1 = tf_pb.paragraphs[0]
p_pb1.text = "Swetha N.  ·  B.E. Computer Science and Design (Class of 2028)"
p_pb1.font.size = Pt(14)
p_pb1.font.bold = True
p_pb1.font.color.rgb = TEXT_PRIMARY
p_pb1.font.name = "Outfit"
p_pb1.alignment = PP_ALIGN.CENTER

p_pb2 = tf_pb.add_paragraph()
p_pb2.text = "Rajalakshmi Engineering College, Chennai  ·  CGPA: 8.65"
p_pb2.font.size = Pt(11)
p_pb2.font.color.rgb = ACCENT_BLUE
p_pb2.font.name = "Inter"
p_pb2.alignment = PP_ALIGN.CENTER
p_pb2.space_before = Pt(3)

# 3 Contact / Resource Cards
connect_cols = [
    ("LIVE DEPLOYMENT", "Vercel Edge Platform", "my-portfolio-...vercel.app", ACCENT_EMERALD),
    ("GITHUB REPOSITORY", "swethanagaraj2007", "github.com/swethanagaraj2007", ACCENT_INDIGO),
    ("PROFESSIONAL CONNECT", "LinkedIn Network", "linkedin.com/in/swetha-n-005697334", ACCENT_BLUE)
]

for i, (tag, val1, val2, col) in enumerate(connect_cols):
    cx = Inches(1.5 + i * 3.55)
    cy = Inches(5.0)
    
    c_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(3.233), Inches(1.35))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = SURFACE_LIGHT
    c_box.line.color.rgb = BORDER_SUBTLE
    c_box.line.width = Pt(1)
    
    # Colored top indicator bar
    c_bar = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(3.233), Inches(0.04))
    c_bar.fill.solid()
    c_bar.fill.fore_color.rgb = col
    c_bar.line.fill.background()
    
    tf_c = c_box.text_frame
    tf_c.word_wrap = True
    
    p_ctag = tf_c.paragraphs[0]
    p_ctag.text = tag
    p_ctag.font.size = Pt(8.5)
    p_ctag.font.bold = True
    p_ctag.font.color.rgb = col
    p_ctag.font.name = "JetBrains Mono"
    
    p_cv1 = tf_c.add_paragraph()
    p_cv1.text = val1
    p_cv1.font.size = Pt(12)
    p_cv1.font.bold = True
    p_cv1.font.color.rgb = TEXT_PRIMARY
    p_cv1.font.name = "Outfit"
    p_cv1.space_before = Pt(4)
    
    p_cv2 = tf_c.add_paragraph()
    p_cv2.text = val2
    p_cv2.font.size = Pt(9.5)
    p_cv2.font.color.rgb = TEXT_MUTED
    p_cv2.font.name = "JetBrains Mono"
    p_cv2.space_before = Pt(2)

# Save the PowerPoint presentation
try:
    prs.save(OUTPUT_PATH)
    print("Redesigned PowerPoint presentation successfully created at:", OUTPUT_PATH)
except PermissionError:
    fallback_path = os.path.join(os.path.dirname(__file__), '..', 'presentation', 'Swetha_Portfolio_Presentation_Latest.pptx')
    prs.save(fallback_path)
    print(f"Original file was open/locked. Saved to fallback path: {fallback_path}")
